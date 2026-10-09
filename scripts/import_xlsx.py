"""Importa o Anexo 2025 sem alterar o arquivo-fonte.

Dependência: openpyxl. O script preserva valores, comentários, espaços, quebras de
linha e a marcação visual da coluna A. Ele nunca deduplica nem corrige conteúdo.
"""

from __future__ import annotations

import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


SOURCE_SHEET = "Anexo 2025"
HEADER_ROW = 11
FIRST_DATA_ROW = 12
LAST_DATA_ROW = 644
EXPECTED_RECORDS = 633

COLUMNS = [
    ("s", "S", "S (marcação da fonte)", False),
    ("sourceNumber", "N°", "N° — coluna B", True),
    ("groupNumber", "N°", "N° — coluna C", True),
    ("law", "LEI", "LEI", True),
    ("exhibitionMode", "n°", "n° — coluna E", True),
    ("onlineExhibitionMode", "ON", "ON", False),
    ("name", "NOME ", "NOME", True),
    ("edition", "EDIÇÃO", "EDIÇÃO", True),
    ("state", "UF", "UF", True),
    ("municipality", "MUNICÍPIO", "MUNICÍPIO", True),
    ("region", "REGIÃO", "REGIÃO", True),
    ("themeProfile", "TEMÁTICA/PERFIL", "TEMÁTICA/PERFIL", True),
    ("scopeProfile", "PERFIL ", "PERFIL", True),
    ("format", "BITOLA", "BITOLA", True),
    ("shortFilmLimit", "ACEITA CURTAS ATÉ", "ACEITA CURTAS ATÉ", True),
    ("registration", "INSCRIÇÃO", "INSCRIÇÃO", True),
    ("registrationFormat", "FORMATO INSCRIÇÃO", "FORMATO INSCRIÇÃO", True),
    ("virtualPlatform", "PLATAFORMA VIRTUAL", "PLATAFORMA VIRTUAL", False),
    ("multiFormat", "MULTIFORMATOS", "MULTIFORMATOS", False),
]

URL_RE = re.compile(r"https?://[^\s)\]}>,;]+", re.I)


def serialize(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def source_fill(cell):
    color = cell.fill.fgColor
    if cell.fill.fill_type is None:
        normalized = None
    elif color.type == "rgb":
        normalized = f"#{str(color.rgb)[-6:]}"
    elif color.type == "theme":
        normalized = f"theme:{color.theme}:tint:{color.tint}"
    elif color.type == "indexed":
        normalized = f"indexed:{color.indexed}"
    else:
        normalized = str(color.type)
    return {"fillType": cell.fill.fill_type, "fillColor": normalized}


def sha256_file(path: Path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalized_name(value):
    return re.sub(r"\s+", " ", value.strip()).casefold()


def main(source_path: str, output_path: str, report_path: str):
    source = Path(source_path)
    output = Path(output_path)
    report_output = Path(report_path)

    workbook = load_workbook(source, data_only=False, keep_links=True)
    values_workbook = load_workbook(source, data_only=True, keep_links=True)
    sheet = workbook[SOURCE_SHEET]
    values_sheet = values_workbook[SOURCE_SHEET]

    actual_headers = [sheet.cell(HEADER_ROW, index).value for index in range(1, 20)]
    expected_headers = [column[1] for column in COLUMNS]
    if actual_headers != expected_headers:
        raise RuntimeError(f"Cabeçalhos divergentes: {actual_headers!r}")

    column_meta = []
    for index, (key, source_header, label, required) in enumerate(COLUMNS, start=1):
        values = [
            serialize(values_sheet.cell(row, index).value)
            for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1)
        ]
        unique = sorted(
            {str(value) for value in values if value not in (None, "")},
            key=str.casefold,
        )
        column_meta.append(
            {
                "key": key,
                "sourceColumn": get_column_letter(index),
                "sourceHeader": source_header,
                "label": label,
                "required": required,
                "distinctValues": unique,
            }
        )

    records = []
    note_count = 0
    records_with_notes = 0
    urls = []

    for source_row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        fields = {}
        notes = {}
        source_cells = {}
        for index, (key, *_rest) in enumerate(COLUMNS, start=1):
            cell = sheet.cell(source_row, index)
            value_cell = values_sheet.cell(source_row, index)
            fields[key] = serialize(value_cell.value)
            source_cells[key] = cell.coordinate
            if cell.comment:
                notes[key] = cell.comment.text
                note_count += 1
                for url in URL_RE.findall(cell.comment.text):
                    urls.append(
                        {
                            "id": f"festival-{int(fields.get('sourceNumber') or source_row):04d}",
                            "field": key,
                            "url": url,
                        }
                    )
        if notes:
            records_with_notes += 1

        source_number = fields["sourceNumber"]
        if not isinstance(source_number, int):
            raise RuntimeError(f"N° inválido na linha {source_row}: {source_number!r}")

        formatting = source_fill(sheet.cell(source_row, 1))
        s_marked = formatting["fillColor"] == "#00B0F0"
        records.append(
            {
                "id": f"festival-{source_number:04d}",
                "sourceRow": source_row,
                "sourceCells": source_cells,
                "fields": fields,
                "notes": notes,
                "sourceFormatting": {"s": formatting},
                "sourceFlags": {"sMarked": s_marked},
            }
        )

    if len(records) != EXPECTED_RECORDS:
        raise RuntimeError(f"Esperados {EXPECTED_RECORDS}, encontrados {len(records)}")
    source_numbers = [record["fields"]["sourceNumber"] for record in records]
    if source_numbers != list(range(1, EXPECTED_RECORDS + 1)):
        raise RuntimeError("A sequência da coluna B deixou de ser 1..633")

    formulas = []
    for row in sheet.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                formula = cell.value
                formulas.append(
                    {
                        "coordinate": cell.coordinate,
                        "formula": getattr(formula, "text", str(formula)),
                        "cachedValue": serialize(values_sheet[cell.coordinate].value),
                    }
                )

    duplicate_names = []
    name_groups = defaultdict(list)
    for record in records:
        name_groups[normalized_name(record["fields"]["name"])].append(record)
    for name, group in name_groups.items():
        if len(group) > 1:
            duplicate_names.append(
                {
                    "normalizedName": name,
                    "ids": [record["id"] for record in group],
                    "editions": [record["fields"]["edition"] for record in group],
                }
            )

    fields_without_numbers = [
        column[0] for column in COLUMNS if column[0] not in {"sourceNumber", "groupNumber"}
    ]
    exact_groups = defaultdict(list)
    for record in records:
        exact_groups[
            tuple(
                normalized_name(str(record["fields"][key]))
                if isinstance(record["fields"][key], str)
                else record["fields"][key]
                for key in fields_without_numbers
            )
        ].append(record["id"])
    exact_duplicates = [ids for ids in exact_groups.values() if len(ids) > 1]

    records_json = json.dumps(records, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    data_hash = hashlib.sha256(records_json.encode("utf-8")).hexdigest()
    source_updated = serialize(values_sheet["J8"].value)

    payload = {
        "meta": {
            "schemaVersion": 1,
            "sourceFile": source.name,
            "sourceSheet": SOURCE_SHEET,
            "sourceUrl": "https://www.panoramadosfestivais.com/textos/2025",
            "sourceUpdatedAt": source_updated,
            "importedAt": "2026-10-08",
            "recordCount": len(records),
            "columnCount": len(COLUMNS),
            "commentCount": note_count,
            "recordsWithNotes": records_with_notes,
            "urlCountInNotes": len(urls),
            "sourceSha256": sha256_file(source),
            "dataSha256": data_hash,
            "columns": column_meta,
            "formulas": formulas,
        },
        "records": records,
    }

    leading_or_trailing = []
    newline_values = []
    for record in records:
        for key, value in record["fields"].items():
            if isinstance(value, str):
                if value != value.strip():
                    leading_or_trailing.append(
                        {"id": record["id"], "field": key, "value": value}
                    )
                if "\n" in value or "\r" in value:
                    newline_values.append(
                        {"id": record["id"], "field": key, "value": value}
                    )

    report = {
        "recordCount": len(records),
        "columnCount": len(COLUMNS),
        "sourceNumber": {
            "min": min(source_numbers),
            "max": max(source_numbers),
            "unique": len(set(source_numbers)),
            "exactSequence": source_numbers == list(range(1, EXPECTED_RECORDS + 1)),
        },
        "comments": {
            "count": note_count,
            "records": records_with_notes,
            "byColumn": dict(
                Counter(
                    get_column_letter(index)
                    for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1)
                    for index in range(1, 20)
                    if sheet.cell(row, index).comment
                )
            ),
        },
        "urlsInNotes": urls,
        "duplicateNameGroups": duplicate_names,
        "exactContentDuplicatesExcludingNumbering": exact_duplicates,
        "sourceFormattingS": dict(
            Counter(
                record["sourceFormatting"]["s"]["fillColor"] or "none"
                for record in records
            )
        ),
        "leadingOrTrailingWhitespace": leading_or_trailing,
        "newlineValues": newline_values,
        "knownSourceInconsistencies": [
            "O autofiltro da planilha termina em G11:S638 e não alcança as linhas 639:644.",
            "O relatório vinculado informa 630 eventos; esta versão atualizada contém 633 registros.",
        ],
    }

    output.parent.mkdir(parents=True, exist_ok=True)
    report_output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    report_output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "records": len(records),
                "columns": len(COLUMNS),
                "comments": note_count,
                "recordsWithNotes": records_with_notes,
                "urls": len(urls),
                "duplicateNameGroups": len(duplicate_names),
                "dataSha256": data_hash,
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    if len(sys.argv) != 4:
        raise SystemExit("uso: import_xlsx.py fonte.xlsx public/data/festivals.json import-report.json")
    main(sys.argv[1], sys.argv[2], sys.argv[3])
